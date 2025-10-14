from typing import Dict, Any, Optional
from pathlib import Path
import re
from docx import Document
from .excel import load_excel
from .locator import Locator
from .renderer import set_paragraph_text_keep_style, replace_underscore_segment_in_paragraph, write_to_cell, replace_between_in_paragraph
from .utils import transform_value, apply_length_policy, normalize
from .renderer import set_paragraph_text_keep_style, replace_underscore_segment_in_paragraph, write_to_cell, replace_between_in_paragraph, replace_after_token_in_paragraph, replace_after_slash_in_paragraph

def apply_mapping_to_doc(doc: Document, cfg_row: Dict[str, Any], map_rows, settings, dry_run: bool) -> list[str]:
    """
    Возвращает список сообщений (для dry-run/лога).
    """
    log: list[str] = []
    min_line_len = int(settings.get("минимальная_длина_линии") or 5)
    policy = str(settings.get("политика_длины_значения") or "underline_and_keep_line")
    locator = Locator(doc, min_line_len=min_line_len)

    for i, rule in enumerate(map_rows, start=1):
        field = (rule.get("Field") or "").strip()
        if field not in cfg_row:
            log.append(f"[{i}] {method} (Field={field}): НЕТ ТАКОЙ КОЛОНКИ В data")
            continue
        method = (rule.get("Method") or "").strip()
        value_raw = cfg_row.get(field)
        # Покажем, какие поля есть и их значения (в dry-run)
        if dry_run:
            present = ", ".join([f"{k}={cfg_row.get(k)!r}" for k in sorted(cfg_row.keys())])
            log.append(f"[DATA] {present}")
        value = transform_value(value_raw if value_raw is not None else rule.get("Default"), rule.get("Transform"))
        occur = int(rule.get("Occur") or 1)

        if value == "":
            log.append(f"[{i}] {method} (Field={field}): ПУСТО — пропуск")
            continue

        if method == "by_number":
            number = int(rule.get("Number") or 0)
            found = locator.by_number_find(number)
            if not found:
                log.append(f"[{i}] by_number #{number}: НЕ НАЙДЕНО (Field={field})")
                continue
            p, seg = found
            log.append(f"[{i}] by_number #{number}: '{seg[:10]}...' -> '{value}'")
            if not dry_run:
                replace_underscore_segment_in_paragraph(p, seg, value, policy)

        elif method == "anchor_after_colon":
            anchor = str(rule.get("Anchor") or "")
            p = locator.anchor_after_colon_find(anchor, occur=occur)
            if not p:
                log.append(f"[{i}] anchor_after_colon '{anchor}' (#{occur}): НЕ НАЙДЕНО (Field={field})")
                continue
            # пишем после двоеточия: заменяем всё после первого ':' на ' ' + value
            text = p.text
            if ":" in text:
                before, after = text.split(":", 1)
                new_text = before + ": " + value
            else:
                new_text = text + " " + value
            log.append(f"[{i}] anchor_after_colon '{anchor}': -> '{value}'")
            if not dry_run:
                set_paragraph_text_keep_style(p, new_text)

        elif method == "anchor_after_slash":
            anchor = str(rule.get("Anchor") or "")
            p = locator.anchor_line_find(anchor, occur=occur)
            if not p:
                log.append(f"[{i}] anchor_after_slash '{anchor}' (#{occur}): НЕ НАЙДЕНО (Field={field})")
                continue
            log.append(f"[{i}] anchor_after_slash '{anchor}' -> '{value}'")
            if not dry_run:
                replace_after_slash_in_paragraph(p, value)

        elif method == "anchor_after_token":
            anchor = str(rule.get("Anchor") or "")
            token = str(rule.get("Label") or "/")
            p = locator.anchor_after_token_find(anchor, token=token, occur=occur)
            if not p:
                log.append(f"[{i}] anchor_after_token '{anchor}' token='{token}' (#{occur}): НЕ НАЙДЕНО (Field={field})")
                continue
            log.append(f"[{i}] anchor_after_token '{anchor}' '{token}': -> '{value}'")
            if not dry_run:
                replace_after_token_in_paragraph(p, token, value)

        elif method == "anchor_segment":
            anchor = str(rule.get("Anchor") or "")
            seg_index = int(rule.get("Segment") or 1)
            found = locator.anchor_segment_find(anchor, seg_index=seg_index, occur=occur)
            if not found:
                log.append(f"[{i}] anchor_segment '{anchor}' seg#{seg_index} (#{occur}): НЕ НАЙДЕНО (Field={field})")
                continue
            p, seg = found
            log.append(f"[{i}] anchor_segment '{anchor}' seg#{seg_index}: '{seg[:10]}...' -> '{value}'")
            if not dry_run:
                replace_underscore_segment_in_paragraph(p, seg, value, policy)

        elif method == "table_label":
            label = str(rule.get("Label") or rule.get("Anchor") or "")
            target = str(rule.get("Target") or "right(1)")
            found = locator.table_label_find(label, target)
            if not found:
                log.append(f"[{i}] table_label '{label}' -> {target}: НЕ НАЙДЕНО (Field={field})")
                continue
            tbl, r, c = found
            log.append(f"[{i}] table_label '{label}' -> {target} [r{r+1},c{c+1}] -> '{value}'")
            if not dry_run:
                write_to_cell(tbl.cell(r, c), value)

        elif method == "cell_ref":
            t_index = int(rule.get("TableIndex") or 1)
            row = int(rule.get("Row") or 1)
            col = int(rule.get("Col") or 1)
            found = locator.cell_ref_find(t_index, row, col)
            if not found:
                log.append(f"[{i}] cell_ref T{t_index}R{row}C{col}: НЕ НАЙДЕНО (Field={field})")
                continue
            tbl, r, c = found
            log.append(f"[{i}] cell_ref T{t_index}R{row}C{col} -> '{value}'")
            if not dry_run:
                write_to_cell(tbl.cell(r, c), value)

        elif method == "anchor_prev":
            anchor = str(rule.get("Anchor") or "")
            p = locator.anchor_prev_find(anchor, occur=occur)
            if not p:
                log.append(f"[{i}] anchor_prev '{anchor}' (#{occur}): НЕ НАЙДЕНО (Field={field})")
                continue
            log.append(f"[{i}] anchor_prev '{anchor}' -> '{value}'")
            if not dry_run:
                set_paragraph_text_keep_style(p, value)

        elif method == "between_words":
            left = str(rule.get("Anchor") or "")
            right = rule.get("Label")  # может быть None/пусто -> до конца строки
            found = locator.between_words_find(left, right if right not in (None, "") else None, occur=occur)
            if not found:
                log.append(f"[{i}] between_words '{left}'..'{right or 'EOL'}' (#{occur}): НЕ НАЙДЕНО (Field={field})")
                continue
            p, lft, rgt = found
            log.append(f"[{i}] between_words '{lft}'..'{rgt or 'EOL'}' -> '{value}'")
            if not dry_run:
                replace_between_in_paragraph(p, lft, rgt if rgt not in (None, "") else None, value)

        else:
            log.append(f"[{i}] НЕИЗВЕСТНЫЙ Method='{method}' (Field={field})")

    return log

def fill_one(doc_path: str, xlsx_path: str, output_path: str, row_index: Optional[int], dry_run: bool) -> int:
    cfg = load_excel(xlsx_path)
    if row_index is not None:
        data_rows = [cfg.data_rows[row_index - 1]]
    else:
        data_rows = cfg.data_rows

    mask = str(cfg.settings.get("маска_имени_файла") or "{{ФИО}}.docx")
    output_dir = Path(output_path)
    output_dir.mkdir(parents=True, exist_ok=True)

    exit_code = 0
    for ridx, row in enumerate(data_rows, start=(row_index or 1)):
        doc = Document(doc_path)
        log = apply_mapping_to_doc(doc, row, cfg.mapping_rows, cfg.settings, dry_run=dry_run)
        # имя файла по маске
        out_name = mask
        for k, v in row.items():
            out_name = out_name.replace(f"{{{{{k}}}}}", str(v or ""))
        out_name = re.sub(r"[\\/:*?\"<>|]", "_", out_name).strip()
        if not out_name.endswith(".docx"):
            out_name += ".docx"
        if dry_run:
            print(f"\n[ROW {ridx}] Предпросмотр '{out_name}':")
            for line in log:
                print("  -", line)
        else:
            out_file = output_dir / out_name
            doc.save(out_file)
            print(f"[ROW {ridx}] Сохранено: {out_file}")
            for line in log:
                print("  -", line)

    return exit_code
